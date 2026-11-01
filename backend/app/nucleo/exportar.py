"""
Exportación de reportes a PDF (fpdf2) y Excel (openpyxl).

Un reporte es una tabla: título, líneas de filtros, columnas (clave, título y tipo) y filas
(diccionarios). Los dos formatos llevan el mismo encabezado: nombre de la unidad educativa,
título, filtros aplicados, fecha y hora de Bolivia y quién lo generó.

Tipos de columna: texto, entero, decimal (2 decimales), porcentaje (número 0-100) y fecha.
"""
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

INSTITUCION = 'Unidad Educativa "Mariscal Sucre G"'
AZUL = (11, 35, 82)          # mismo azul del menú lateral del sistema
GRIS = (241, 245, 249)


@dataclass
class Columna:
    clave: str
    titulo: str
    tipo: str = "texto"


@dataclass
class Tabla:
    titulo: str
    columnas: list[Columna]
    filas: list[dict]
    filtros: list[str] = field(default_factory=list)
    resumen: list[str] = field(default_factory=list)   # líneas al pie (totales, promedios)
    generado_por: str = ""
    generado_en: datetime | None = None


def texto_celda(valor, tipo: str) -> str:
    """Valor como se muestra en pantalla y en el PDF."""
    if valor is None or valor == "":
        return "-"
    if tipo == "fecha" and isinstance(valor, (date, datetime)):
        return valor.strftime("%d/%m/%Y")
    if tipo == "decimal":
        return f"{float(valor):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if tipo == "porcentaje":
        return f"{float(valor):.1f} %".replace(".", ",")
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    return str(valor)


def _latin1(t: str) -> str:
    """Las fuentes base del PDF son latin-1: se cambian los pocos signos que no entran."""
    t = t.replace("—", "-").replace("–", "-").replace("“", '"').replace("”", '"').replace("·", "-").replace("…", "...")
    return t.encode("latin-1", "replace").decode("latin-1")


def _pie(tabla: Tabla) -> str:
    cuando = (tabla.generado_en or datetime.now()).strftime("%d/%m/%Y %H:%M")
    quien = f" por {tabla.generado_por}" if tabla.generado_por else ""
    return f"Generado el {cuando} (hora de Bolivia){quien}"


# ------------------------------------------------------------------ PDF

class _Pdf(FPDF):
    def __init__(self, tabla: Tabla, orientacion: str):
        super().__init__(orientation=orientacion, unit="mm", format="Letter")
        self.tabla = tabla
        self.set_margins(12, 12, 12)
        self.set_auto_page_break(True, margin=14)

    def header(self):
        self.set_font("helvetica", "B", 9)
        self.set_text_color(*AZUL)
        self.cell(0, 5, _latin1(INSTITUCION), new_x="LMARGIN", new_y="NEXT")
        self.set_font("helvetica", "B", 13)
        self.set_text_color(15, 23, 42)
        self.cell(0, 7, _latin1(self.tabla.titulo), new_x="LMARGIN", new_y="NEXT")
        self.set_font("helvetica", "", 8.5)
        self.set_text_color(71, 85, 105)
        for f in self.tabla.filtros:
            self.cell(0, 4.5, _latin1(f), new_x="LMARGIN", new_y="NEXT")
        self.ln(2)
        self.set_text_color(0, 0, 0)

    def footer(self):
        self.set_y(-10)
        self.set_font("helvetica", "", 7.5)
        self.set_text_color(100, 116, 139)
        self.cell(0, 5, _latin1(_pie(self.tabla)), align="L")
        self.cell(0, 5, f"Página {self.page_no()} de {{nb}}".encode("latin-1").decode("latin-1"), align="R")


def a_pdf(tabla: Tabla) -> bytes:
    columnas = tabla.columnas
    # Ancho relativo de cada columna según su contenido (con tope para que los textos largos bajen de línea)
    pesos = []
    for c in columnas:
        largo = max([len(c.titulo)] + [len(texto_celda(f.get(c.clave), c.tipo)) for f in tabla.filas[:300]])
        pesos.append(min(max(largo, 4), 40))
    horizontal = sum(pesos) > 110 or len(columnas) > 7
    pdf = _Pdf(tabla, "L" if horizontal else "P")
    pdf.alias_nb_pages()
    pdf.add_page()
    letra = 8 if len(columnas) <= 8 else 7
    pdf.set_font("helvetica", "", letra)
    alineacion = tuple("LEFT" if c.tipo in ("texto", "fecha") else "RIGHT" for c in columnas)
    encabezado = FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=AZUL)
    with pdf.table(col_widths=tuple(pesos), text_align=alineacion, headings_style=encabezado, repeat_headings=1,
                   cell_fill_color=GRIS, cell_fill_mode="ROWS", line_height=4.6, padding=1.2,
                   borders_layout="HORIZONTAL_LINES") as t:
        fila = t.row()
        for c in columnas:
            fila.cell(_latin1(c.titulo))
        for f in tabla.filas:
            fila = t.row()
            for c in columnas:
                fila.cell(_latin1(texto_celda(f.get(c.clave), c.tipo)))
    if not tabla.filas:
        pdf.ln(3)
        pdf.cell(0, 6, _latin1("No hay datos para los filtros elegidos."), new_x="LMARGIN", new_y="NEXT")
    if tabla.resumen:
        pdf.ln(3)
        pdf.set_font("helvetica", "B", 8.5)
        for r in tabla.resumen:
            pdf.cell(0, 5, _latin1(r), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


# ------------------------------------------------------------------ Excel

def _valor_excel(valor, tipo: str):
    if valor is None:
        return None
    if isinstance(valor, Decimal):
        return float(valor)
    if tipo == "fecha" and isinstance(valor, datetime):
        return valor.replace(tzinfo=None)
    return valor


def a_excel(tabla: Tabla) -> bytes:
    libro = Workbook()
    hoja = libro.active
    # Excel no admite estos caracteres en el nombre de la hoja (máximo 31)
    hoja.title = "".join("-" if ch in "[]:*?/\\" else ch for ch in tabla.titulo)[:31].strip()
    hoja["A1"] = INSTITUCION
    hoja["A1"].font = Font(bold=True, color="0B2352")
    hoja["A2"] = tabla.titulo
    hoja["A2"].font = Font(bold=True, size=14)
    fila = 3
    for f in tabla.filtros + [_pie(tabla)]:
        hoja.cell(row=fila, column=1, value=f).font = Font(color="475569", size=9)
        fila += 1
    fila += 1
    inicio = fila
    for i, c in enumerate(tabla.columnas, start=1):
        celda = hoja.cell(row=fila, column=i, value=c.titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="0B2352")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    formatos = {"decimal": "#,##0.00", "porcentaje": "0.0", "fecha": "DD/MM/YYYY", "entero": "0"}
    for f in tabla.filas:
        fila += 1
        for i, c in enumerate(tabla.columnas, start=1):
            celda = hoja.cell(row=fila, column=i, value=_valor_excel(f.get(c.clave), c.tipo))
            if c.tipo in formatos:
                celda.number_format = formatos[c.tipo]
    ultima = get_column_letter(len(tabla.columnas))
    if tabla.filas:
        hoja.auto_filter.ref = f"A{inicio}:{ultima}{fila}"
    hoja.freeze_panes = f"A{inicio + 1}"
    for i, c in enumerate(tabla.columnas, start=1):
        largo = max([len(c.titulo)] + [len(texto_celda(f.get(c.clave), c.tipo)) for f in tabla.filas[:500]])
        hoja.column_dimensions[get_column_letter(i)].width = min(max(largo + 2, 8), 45)
    for r in tabla.resumen:
        fila += 2 if r is tabla.resumen[0] else 1
        hoja.cell(row=fila, column=1, value=r).font = Font(bold=True)
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()

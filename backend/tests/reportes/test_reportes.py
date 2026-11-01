"""Pruebas del módulo de reportes (estáticos y dinámico, en JSON, PDF y Excel).

Datos de 08_datos_prueba: gestión 2026 (id 1), curso 1ro A - Secundaria (id 14) con 6 inscritos,
1er trimestre (id 1) con notas y asistencia, y cuotas atrasadas.
"""
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import text

from app.nucleo import bd
from tests.conftest import USUARIOS, iniciar


async def _valor(consulta: str, **params):
    async with bd._fabrica() as sesion:
        return await sesion.scalar(text(consulta), params)


async def _descargas():
    return await _valor("SELECT count(*) FROM bitacora WHERE modulo = 'REPORTES'")


async def test_reportes_solo_administrador(cliente, admin):
    r = await cliente.get("/reportes/estaticos", headers=admin)
    assert r.status_code == 200 and [x["clave"] for x in r.json()] == ["morosos", "inscritos", "asistencia", "rendimiento"]
    for usuario in ("profesor", "tutor", "estudiante"):
        cabecera = await iniciar(cliente, USUARIOS[usuario])
        assert (await cliente.get("/reportes/estaticos", headers=cabecera)).status_code == 403
        assert (await cliente.post("/reportes/dinamico", headers=cabecera,
                                   json={"fuente": "pagos", "columnas": ["ci"]})).status_code == 403


async def test_reportes_opciones(cliente, admin):
    r = await cliente.get("/reportes/opciones", headers=admin)
    assert r.status_code == 200
    datos = r.json()
    assert datos["gestiones"] and datos["cursos"] and datos["periodos"] and datos["materias"]


async def test_morosos_en_los_tres_formatos(cliente, admin):
    r = await cliente.get("/reportes/estaticos/morosos", headers=admin)
    assert r.status_code == 200, r.text
    datos = r.json()
    assert datos["titulo"] == "Reporte de morosos" and datos["filtros"] == ["Gestión: 2026"]
    esperado = await _valor("SELECT count(*) FROM v_estado_pagos WHERE estado_actual = 'atrasado' AND anio = 2026")
    assert datos["total"] == esperado and "Total adeudado" in datos["resumen"][0]
    assert isinstance(datos["filas"][0]["monto"], float)

    antes = await _descargas()
    r = await cliente.get("/reportes/estaticos/morosos", headers=admin, params={"formato": "pdf"})
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF") and "attachment" in r.headers["content-disposition"]
    r = await cliente.get("/reportes/estaticos/morosos", headers=admin, params={"formato": "xlsx"})
    assert r.status_code == 200 and r.content.startswith(b"PK")
    hoja = load_workbook(BytesIO(r.content)).active
    assert hoja["A2"].value == "Reporte de morosos"
    assert await _descargas() == antes + 2


async def test_inscritos_por_curso(cliente, admin):
    r = await cliente.get("/reportes/estaticos/inscritos", headers=admin, params={"id_gestion": 1})
    assert r.status_code == 400 and r.json()["detalle"] == "Elija el curso."
    r = await cliente.get("/reportes/estaticos/inscritos", headers=admin, params={"id_gestion": 1, "id_curso": 14})
    assert r.status_code == 200, r.text
    datos = r.json()
    assert datos["total"] == 6 and "Curso: 1ro A - Secundaria" in datos["filtros"]
    assert [f["n"] for f in datos["filas"]] == list(range(1, 7))


async def test_asistencia_y_rendimiento_por_curso(cliente, admin):
    for clave in ("asistencia", "rendimiento"):
        r = await cliente.get(f"/reportes/estaticos/{clave}", headers=admin, params={"id_curso": 14, "id_periodo": 1})
        assert r.status_code == 200, r.text
        datos = r.json()
        assert datos["filas"] and datos["resumen"]
    situaciones = {f["situacion"] for f in datos["filas"]}
    assert situaciones <= {"Aprobado", "Reprobado", "Sin notas"}
    r = await cliente.get("/reportes/estaticos/rendimiento", headers=admin, params={"id_curso": 14})
    assert r.status_code == 400
    r = await cliente.get("/reportes/estaticos/inexistente", headers=admin)
    assert r.status_code == 404


async def test_dinamico_filtra_y_cuenta(cliente, admin):
    r = await cliente.get("/reportes/dinamico/fuentes", headers=admin)
    assert [f["clave"] for f in r.json()] == ["estudiantes", "pagos", "asistencia", "calificaciones"]
    cuerpo = {"fuente": "estudiantes", "columnas": ["ci", "estudiante", "curso"], "filtros": {"curso": 14}}
    r = await cliente.post("/reportes/dinamico", headers=admin, json=cuerpo)
    assert r.status_code == 200, r.text
    datos = r.json()
    assert datos["total"] == 6 and [c["clave"] for c in datos["columnas"]] == ["ci", "estudiante", "curso"]
    assert all(f["curso"] == "1ro A - Secundaria" for f in datos["filas"])
    assert "Curso: 1ro A - Secundaria" in datos["filtros"]


async def test_dinamico_ordena(cliente, admin):
    cuerpo = {"fuente": "calificaciones", "columnas": ["estudiante", "materia", "nota_final"],
              "filtros": {"periodo": 1, "curso": 14}, "orden": "nota_final", "descendente": True}
    notas = [f["nota_final"] for f in (await cliente.post("/reportes/dinamico", headers=admin, json=cuerpo)).json()["filas"]]
    assert notas and notas == sorted(notas, reverse=True)


async def test_dinamico_rechaza_lo_que_no_esta_en_la_lista(cliente, admin):
    malos = [
        {"fuente": "usuarios", "columnas": ["ci"]},
        {"fuente": "pagos", "columnas": ["ci; DROP TABLE persona"]},
        {"fuente": "pagos", "columnas": ["password_hash"]},
        {"fuente": "pagos", "columnas": ["ci"], "filtros": {"otro": 1}},
        {"fuente": "pagos", "columnas": ["ci"], "filtros": {"estado": "cualquiera"}},
        {"fuente": "pagos", "columnas": ["ci"], "filtros": {"pago_desde": "ayer"}},
        {"fuente": "pagos", "columnas": ["ci"], "filtros": {"curso": "1 OR 1=1"}},
        {"fuente": "pagos", "columnas": ["ci"], "orden": "1; DELETE"},
    ]
    for cuerpo in malos:
        r = await cliente.post("/reportes/dinamico", headers=admin, json=cuerpo)
        assert r.status_code == 400, (cuerpo, r.text)
    assert await _valor("SELECT count(*) FROM persona") > 0


async def test_dinamico_en_pdf_y_excel(cliente, admin):
    cuerpo = {"fuente": "pagos", "columnas": ["estudiante", "mes", "monto", "estado"], "filtros": {"estado": "atrasado"}}
    antes = await _descargas()
    r = await cliente.post("/reportes/dinamico", headers=admin, json=cuerpo, params={"formato": "pdf"})
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    r = await cliente.post("/reportes/dinamico", headers=admin, json=cuerpo, params={"formato": "xlsx"})
    hoja = load_workbook(BytesIO(r.content)).active
    assert hoja["A2"].value == "Reporte personalizado: Pagos de mensualidad"
    assert await _descargas() == antes + 2
    # La vista previa no se anota en la bitácora
    await cliente.post("/reportes/dinamico", headers=admin, json=cuerpo)
    assert await _descargas() == antes + 2
